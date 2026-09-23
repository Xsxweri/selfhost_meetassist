import io
import json
from app.models.meeting import Meeting
from app.models.action_item import ActionItem
from app.models.transcript import Transcript


# === 导出服务 ===
def _esc(s: str | None) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ===== 辅助函数 =====
def _item_meta(a: ActionItem) -> str:
    meta = []
    if a.assignee:
        meta.append(f"负责人：{a.assignee}")
    if a.due_date:
        meta.append(f"截止：{a.due_date:%Y-%m-%d}")
    meta.append(f"优先级：{a.priority}")
    meta.append(f"状态：{a.status}")
    return "，".join(meta)


# ======= 导出为 Markdown =======
def to_markdown(meeting: Meeting, items: list[ActionItem], transcripts: list[Transcript]) -> str:
    md = [f"# {meeting.title}", ""]
    md.append(f"- 创建时间：{meeting.created_at:%Y-%m-%d %H:%M}")
    md.append("")
    md.append("## 会议纪要")
    md.append(meeting.summary or "（暂无纪要）")
    md.append("")
    md.append("## 待办事项")
    if items:
        for a in items:
            md.append(f"- [ ] {a.content}（{_item_meta(a)}）")
    else:
        md.append("（无）")
    md.append("")
    md.append("## 转录全文")
    if transcripts:
        for t in transcripts:
            if t.text:
                md.append(t.text)
    else:
        md.append("（暂无转录）")
    return "\n".join(md)


# ======= 导出为 JSON =======
def to_json(meeting: Meeting, items: list[ActionItem], transcripts: list[Transcript]) -> str:
    data = {
        "id": str(meeting.id),
        "title": meeting.title,
        "created_at": meeting.created_at.isoformat(),
        "summary": meeting.summary,
        "action_items": [
            {
                "content": a.content,
                "assignee": a.assignee,
                "due_date": a.due_date.isoformat() if a.due_date else None,
                "priority": a.priority,
                "status": a.status,
            }
            for a in items
        ],
        "transcripts": [t.text for t in transcripts if t.text],
    }
    return json.dumps(data, ensure_ascii=False, indent=2)


# ======= 导出为 DOCX =======
def to_docx(meeting: Meeting, items: list[ActionItem], transcripts: list[Transcript]) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(meeting.title, level=0)
    doc.add_heading("会议纪要", level=1)
    doc.add_paragraph(meeting.summary or "（暂无纪要）")
    doc.add_heading("待办事项", level=1)
    if items:
        for a in items:
            doc.add_paragraph(f"{a.content}（{_item_meta(a)}）", style="List Bullet")
    else:
        doc.add_paragraph("（无）")
    doc.add_heading("转录全文", level=1)
    for t in transcripts:
        if t.text:
            doc.add_paragraph(t.text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ======= 导出为 PDF =======
def to_pdf(meeting: Meeting, items: list[ActionItem], transcripts: list[Transcript]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    base = getSampleStyleSheet()["Normal"]
    zh = ParagraphStyle("zh", parent=base, fontName="STSong-Light", fontSize=10.5, leading=16)
    h1 = ParagraphStyle("h1", parent=zh, fontSize=18, leading=24, spaceAfter=10)
    h2 = ParagraphStyle("h2", parent=zh, fontSize=14, leading=20, spaceBefore=10, spaceAfter=6)

    story = [Paragraph(_esc(meeting.title), h1)]
    story.append(Paragraph("会议纪要", h2))
    story.append(Paragraph(_esc(meeting.summary or "（暂无纪要）").replace("\n", "<br/>"), zh))
    story.append(Paragraph("待办事项", h2))
    if items:
        for a in items:
            story.append(Paragraph("• " + _esc(f"{a.content}（{_item_meta(a)}）"), zh))
    else:
        story.append(Paragraph("（无）", zh))
    story.append(Paragraph("转录全文", h2))
    for t in transcripts:
        if t.text:
            story.append(Paragraph(_esc(t.text), zh))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
    )
    doc.build(story)
    return buf.getvalue()


def build(fmt: str, meeting: Meeting, items: list[ActionItem], transcripts: list[Transcript]):
    """统一出口：返回 (content_bytes, media_type, ext)"""
    if fmt == "md":
        return to_markdown(meeting, items, transcripts).encode("utf-8"), "text/markdown; charset=utf-8", ".md"
    if fmt == "json":
        return to_json(meeting, items, transcripts).encode("utf-8"), "application/json; charset=utf-8", ".json"
    if fmt == "docx":
        return to_docx(meeting, items, transcripts), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"
    if fmt == "pdf":
        return to_pdf(meeting, items, transcripts), "application/pdf", ".pdf"
    raise ValueError(f"unsupported format: {fmt}")